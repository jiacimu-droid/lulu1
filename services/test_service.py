import json
import os
import tempfile
import shutil
import unittest
from pathlib import Path
from unittest.mock import patch
import lulu_service as service

class DurableServiceTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_data = service.DATA
        service.DATA = Path(self.temp.name)
        service.initialize()
    def tearDown(self):
        service.DATA = self.old_data
        self.temp.cleanup()
    def test_duplicate_submission_does_not_duplicate_job(self):
        a = service.submit('a', 'request', 'docx', {'notes': 'real notes'})
        b = service.submit('a', 'request', 'docx', {'notes': 'real notes'})
        self.assertEqual(a['id'], b['id'])
        with self.assertRaises(ValueError): service.submit('a', 'request', 'pptx', {'notes': 'changed'})
    def test_character_tasks_do_not_share_identity(self):
        a = service.submit('a', 'request', 'docx', {'notes': 'A'})
        b = service.submit('b', 'request', 'docx', {'notes': 'B'})
        self.assertNotEqual(a['id'], b['id'])
    def test_expired_lease_recovers_after_restart(self):
        service.submit('a', 'request', 'docx', {'notes': 'real notes'})
        first = service.claim(now=100)
        service.initialize()  # Reopen the existing persistent database.
        self.assertIsNone(service.claim(now=200))
        recovered = service.claim(now=1001)
        self.assertEqual(first['id'], recovered['id'])
        self.assertEqual(recovered['attempt'], 2)
    def test_repeated_worker_death_ends_with_truthful_failure(self):
        service.submit('a', 'request', 'docx', {'notes': 'real notes'})
        for t in (100, 1001, 2002): self.assertIsNotNone(service.claim(now=t))
        self.assertIsNone(service.claim(now=3003))
        with service.connection() as db:
            self.assertEqual(db.execute('SELECT status FROM jobs').fetchone()[0], 'failed')
    def test_claude_conversion_uses_native_tools_and_phone_context(self):
        messages = [{'role': 'system', 'content': 'replace persona'}, {'role': 'user', 'content': '闹钟'},
                    {'role': 'assistant', 'content': None, 'tool_calls': [{'id': 'tool1', 'function': {'name': 'create_alarm', 'arguments': '{"time":"10:00"}'}}]},
                    {'role': 'tool', 'tool_call_id': 'tool1', 'content': '{"success":true}'}]
        with patch.dict(os.environ, {'CLAUDE_MODEL': 'configured-test-model'}):
            body = service.claude_payload('locked phone core', messages,
                [{'function': {'name': 'create_alarm', 'parameters': {'type': 'object'}}}])
        self.assertEqual(body['system'], 'locked phone core')
        self.assertEqual(body['messages'][1]['content'][0]['type'], 'tool_use')
        self.assertEqual(body['messages'][2]['content'][0]['type'], 'tool_result')
        self.assertIn('input_schema', body['tools'][0])
    def test_voice_agent_must_use_shared_core(self):
        config = {'conversation_config': {'agent': {'first_message': '', 'prompt': {'llm': 'custom-llm', 'custom_llm': {'url': 'https://lulu.example/v1/llm/chat/completions'}}}}}
        service.validate_voice_agent(config, 'https://lulu.example')
        config['conversation_config']['agent']['prompt']['llm'] = 'other'
        with self.assertRaises(ValueError): service.validate_voice_agent(config, 'https://lulu.example')
    def test_research_cannot_claim_notes_are_retrieved_sources(self):
        job = service.submit('a', 'research-request', 'research', {'notes': 'only notes'})
        with patch('lulu_service.sources_for', return_value=[]):
            with self.assertRaises(ValueError): service.execute_job(job)
    def test_resumed_job_uses_saved_sources(self):
        job = service.submit('a', 'sources-request', 'research', {'sources':['https://example.org']})
        directory = service.DATA / 'artifacts' / job['id']
        directory.mkdir(parents=True)
        sources = [{'url':'https://example.org','text':'saved source'}]
        (directory / 'sources.json').write_text(json.dumps(sources))
        with patch('lulu_service.sources_for', side_effect=AssertionError('must not refetch')):
            result = service.execute_job(job)
        self.assertEqual(result['sources'], ['https://example.org'])
    def test_model_name_is_required_not_invented(self):
        with patch.dict(os.environ, {'CLAUDE_MODEL': ''}):
            with self.assertRaises(ValueError): service.claude_payload('core', [])
    def test_unverified_reference_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'Unverified citation'):
            service.build_artifact({'kind':'docx'}, {'title':'test','sections':[{'title':'x','citations':[7]}]}, [], service.DATA)
    @unittest.skipUnless(shutil.which('libreoffice'), 'LibreOffice unavailable in this runtime')
    def test_real_office_files_and_rendered_previews(self):
        from pypdf import PdfReader
        from pptx import Presentation
        from docx import Document
        outline = {'title': 'Lulu artifact verification', 'sections': [{'title':'Evidence', 'paragraphs':['Persisted real content'], 'citations':[1]}]}
        sources = [{'url':'https://example.org/provided-source', 'text':'user supplied test fixture'}]
        for kind in ('docx', 'pptx'):
            with self.subTest(kind=kind):
                directory = service.DATA / kind
                directory.mkdir()
                result = service.build_artifact({'kind':kind}, outline, sources, directory)
                self.assertEqual(result['sha256'], __import__('hashlib').sha256((directory/result['file']).read_bytes()).hexdigest())
                preview = PdfReader(directory/result['preview'])
                text = '\n'.join(page.extract_text() for page in preview.pages)
                self.assertIn('Persisted real content', text)
                if kind == 'pptx': self.assertEqual(len(Presentation(directory/result['file']).slides), 3)
                else: self.assertTrue(Document(directory/result['file']).paragraphs)
    def test_missing_renderer_cannot_claim_office_success(self):
        with patch('lulu_service.subprocess.run', side_effect=FileNotFoundError('libreoffice')):
            with self.assertRaises(FileNotFoundError):
                service.build_artifact({'kind':'docx'}, {'title':'test','sections':[{'title':'x','paragraphs':['真实原文']} ]}, [], service.DATA)
        self.assertTrue((service.DATA/'result.docx').is_file())
        from docx import Document
        self.assertIn('真实原文', [p.text for p in Document(service.DATA/'result.docx').paragraphs])

if __name__ == '__main__': unittest.main()
